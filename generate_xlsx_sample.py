"""
Generates a sample Excel workbook (xlsx_samples/sample.xlsx) for manually
testing Log Guard's .xlsx support, and can print the contents of any .xlsx
file so you can see what was (or wasn't) masked — no Excel needed.

Usage:
    python generate_xlsx_sample.py               # create xlsx_samples/sample.xlsx
    python generate_xlsx_sample.py --show FILE   # print every non-empty cell in FILE
"""
import os
import sys
from datetime import datetime

import openpyxl

OUTPUT_DIR = "xlsx_samples"
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "sample.xlsx")


def create_sample_workbook(filepath):
    """
    Builds a 4-sheet workbook full of fake sensitive data, covering:
    text cells, a phone number stored as a number, dates, a decimal,
    a formula (which should be left alone), and a hidden sheet.
    """
    workbook = openpyxl.Workbook()

    users = workbook.active
    users.title = "Users"
    users.append(["Name", "Email", "Phone", "Signup date", "Notes"])
    users.append(["Alice", "alice@example.com", 5551234567, datetime(2024, 1, 15), "Called from 555-987-6543"])
    users.append(["Bob", "bob.smith@company.org", "(555) 222-3333", datetime(2024, 2, 20), "No issues"])
    users["A4"] = "Support"
    users["B4"] = '=HYPERLINK("mailto:support@example.com","Email support")'

    server = workbook.create_sheet("Server")
    server.append(["Timestamp", "Message"])
    server.append(["2024-01-15 10:22:31", "Login from 192.168.1.10"])
    server.append(["2024-01-15 10:23:02", "API request with key sk_live_51Hz8f92jak3ndlka9d"])
    server.append([3.75, "Payment card 4111 1111 1111 1111 charged"])

    archive = workbook.create_sheet("Archive")
    archive["A1"] = "Old contact: carol@example.net"
    archive["A2"] = "Backup server 10.0.0.5"
    archive.sheet_state = "hidden"

    summary = workbook.create_sheet("Summary")
    summary["A1"] = "Report generated"
    summary["B1"] = 42

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    workbook.save(filepath)


def show_workbook(filepath):
    """
    Prints every non-empty cell of every sheet (hidden sheets included),
    one per line, as  CELL: value
    """
    workbook = openpyxl.load_workbook(filepath)
    for worksheet in workbook.worksheets:
        state = "" if worksheet.sheet_state == "visible" else f" ({worksheet.sheet_state})"
        print(f"=== Sheet: {worksheet.title}{state} ===")
        for row in worksheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    print(f"  {cell.coordinate}: {cell.value!r}")


def main():
    if len(sys.argv) == 3 and sys.argv[1] == "--show":
        show_workbook(sys.argv[2])
    elif len(sys.argv) == 1:
        create_sample_workbook(OUTPUT_FILE)
        print(f"Created {OUTPUT_FILE}")
    else:
        print(__doc__)
        sys.exit(2)


if __name__ == "__main__":
    main()