import sys
import pytest
from datetime import datetime

import openpyxl

from log_guard import (
    mask_line,
    mask_log,
    read_log_file_lines,
    build_output_filepath,
    build_patterns,
    process_file,
    merge_counts,
    expand_filepaths,
    load_config,
    main,
    mask_workbook,
    DEFAULT_EXTENSIONS,
)


def test_mask_line_email():
    patterns = build_patterns()
    line = "User login successful for john.doe@example.com\n"
    masked, counts = mask_line(line, patterns)
    assert "[EMAIL_REDACTED]" in masked
    assert counts["email"] == 1


def test_mask_line_phone():
    patterns = build_patterns()
    line = "Failed login attempt, contact support at 555-123-4567\n"
    masked, counts = mask_line(line, patterns)
    assert "[PHONE_REDACTED]" in masked
    assert counts["phone"] == 1


def test_mask_line_api_key():
    patterns = build_patterns()
    line = "API request with key sk_live_51Hz8f92jak3ndlka9d\n"
    masked, counts = mask_line(line, patterns)
    assert "[API_KEY_REDACTED]" in masked
    assert counts["api_key"] == 1


def test_mask_line_ip_address():
    patterns = build_patterns()
    line = "Connection from 192.168.1.100 accepted\n"
    masked, counts = mask_line(line, patterns)
    assert "[IP_ADDRESS_REDACTED]" in masked
    assert counts["ip_address"] == 1


def test_mask_line_credit_card():
    patterns = build_patterns()
    line = "Payment made with card 4111 1111 1111 1111\n"
    masked, counts = mask_line(line, patterns)
    assert "[CREDIT_CARD_REDACTED]" in masked
    assert counts["credit_card"] == 1


def test_mask_line_multiple_types_in_one_line():
    patterns = build_patterns()
    line = "Contact john.doe@example.com or call 555-123-4567\n"
    masked, counts = mask_line(line, patterns)
    assert "[EMAIL_REDACTED]" in masked
    assert "[PHONE_REDACTED]" in masked
    assert counts["email"] == 1
    assert counts["phone"] == 1


def test_mask_line_no_sensitive_data():
    patterns = build_patterns()
    line = "This is a normal log line with nothing sensitive.\n"
    masked, counts = mask_line(line, patterns)
    assert masked == line
    assert all(count == 0 for count in counts.values())


def test_build_output_filepath():
    assert build_output_filepath("sample.log") == "sample.masked.log"
    assert build_output_filepath("server.txt") == "server.masked.txt"


def test_merge_counts():
    a = {"email": 2, "phone": 0}
    b = {"email": 1, "phone": 3}
    assert merge_counts(a, b) == {"email": 3, "phone": 3}


def test_process_file_normal_mode_writes_output_file(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "test.log"
    log_file.write_text("Contact john.doe@example.com\n", encoding="utf-8")

    result = process_file(str(log_file), patterns, check_mode=False)

    assert result["error"] is None
    assert result["counts"]["email"] == 1
    output_path = tmp_path / "test.masked.log"
    assert output_path.exists()
    assert "[EMAIL_REDACTED]" in output_path.read_text(encoding="utf-8")


def test_process_file_check_mode_writes_no_output_file(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "test.log"
    log_file.write_text("Contact john.doe@example.com\n", encoding="utf-8")

    result = process_file(str(log_file), patterns, check_mode=True)

    assert result["output_filepath"] is None
    output_path = tmp_path / "test.masked.log"
    assert not output_path.exists()


def test_process_file_missing_file():
    patterns = build_patterns()
    result = process_file("does_not_exist.log", patterns, check_mode=True)
    assert result["error"] is not None
    assert "file not found" in result["error"]


def test_process_file_directory_instead_of_file(tmp_path):
    patterns = build_patterns()
    result = process_file(str(tmp_path), patterns, check_mode=True)
    assert result["error"] is not None
    assert "file not found" in result["error"]


def test_process_file_empty_file(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "empty.log"
    log_file.write_text("", encoding="utf-8")

    result = process_file(str(log_file), patterns, check_mode=True)
    assert result["empty"] is True
    assert result["line_count"] == 0


def test_process_file_skips_output_when_clean_and_no_override(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "clean.log"
    log_file.write_text("Nothing sensitive on this line.\n", encoding="utf-8")

    result = process_file(str(log_file), patterns, check_mode=False)

    assert result["skipped_write"] is True
    assert result["output_filepath"] is None
    output_path = tmp_path / "clean.masked.log"
    assert not output_path.exists()


def test_process_file_writes_output_when_clean_but_override_given(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "clean.log"
    log_file.write_text("Nothing sensitive on this line.\n", encoding="utf-8")
    override_path = tmp_path / "custom_output.log"

    result = process_file(str(log_file), patterns, check_mode=False, output_override=str(override_path))

    assert result["skipped_write"] is False
    assert result["output_filepath"] == str(override_path)
    assert override_path.exists()


def test_process_file_output_override_used_for_dirty_file(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "test.log"
    log_file.write_text("Contact john.doe@example.com\n", encoding="utf-8")
    override_path = tmp_path / "custom_output.log"

    result = process_file(str(log_file), patterns, check_mode=False, output_override=str(override_path))

    assert result["output_filepath"] == str(override_path)
    assert override_path.exists()
    default_path = tmp_path / "test.masked.log"
    assert not default_path.exists()


def test_full_masking_workflow_with_sample_data(tmp_path):
    patterns = build_patterns()
    sample_content = (
        "User login successful for john.doe@example.com\n"
        "API request with key sk_live_51Hz8f92jak3ndlka9d\n"
        "Failed login attempt, contact support at 555-123-4567\n"
    )
    log_file = tmp_path / "test.log"
    log_file.write_text(sample_content, encoding="utf-8")

    lines = read_log_file_lines(str(log_file))
    masked_lines, total_counts, line_count = mask_log(lines, patterns)

    assert line_count == 3
    assert total_counts["email"] == 1
    assert total_counts["phone"] == 1
    assert total_counts["api_key"] == 1
    assert "john.doe@example.com" not in "".join(masked_lines)


def test_file_with_bad_encoding_does_not_crash(tmp_path):
    log_file = tmp_path / "bad_encoding.log"
    log_file.write_bytes(b"Normal line\nBroken line: \xff\xfe end\n")

    lines = list(read_log_file_lines(str(log_file)))
    assert len(lines) == 2


def test_custom_pattern_from_json(tmp_path):
    patterns_file = tmp_path / "custom_patterns.json"
    patterns_file.write_text(r'{"employee_id": "EMP-\\d{6}"}', encoding="utf-8")

    patterns = build_patterns(str(patterns_file))
    assert "employee_id" in patterns
    assert "email" in patterns

    line = "Access granted for employee EMP-482913\n"
    masked, counts = mask_line(line, patterns)
    assert "[EMPLOYEE_ID_REDACTED]" in masked
    assert counts["employee_id"] == 1


def test_custom_pattern_invalid_regex_raises_clear_error(tmp_path):
    patterns_file = tmp_path / "bad_regex.json"
    patterns_file.write_text('{"broken": "["}', encoding="utf-8")

    with pytest.raises(ValueError) as exc_info:
        build_patterns(str(patterns_file))
    assert "broken" in str(exc_info.value)


def test_custom_pattern_not_a_dict_raises_clear_error(tmp_path):
    patterns_file = tmp_path / "not_a_dict.json"
    patterns_file.write_text('["email", "phone"]', encoding="utf-8")

    with pytest.raises(ValueError):
        build_patterns(str(patterns_file))


def test_custom_pattern_non_string_value_raises_clear_error(tmp_path):
    patterns_file = tmp_path / "bad_value.json"
    patterns_file.write_text('{"employee_id": 12345}', encoding="utf-8")

    with pytest.raises(ValueError):
        build_patterns(str(patterns_file))


def test_line_with_no_trailing_newline(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "no_newline.log"
    log_file.write_text("Contact john.doe@example.com", encoding="utf-8")

    lines = list(read_log_file_lines(str(log_file)))
    assert len(lines) == 1
    masked_lines, total_counts, line_count = mask_log(lines, patterns)
    assert line_count == 1
    assert total_counts["email"] == 1


def test_line_with_unicode_and_emoji_does_not_crash():
    patterns = build_patterns()
    line = "User \U0001F600 login from user@example.com successful \U0001F389\n"
    masked, counts = mask_line(line, patterns)
    assert counts["email"] == 1
    assert "[EMAIL_REDACTED]" in masked
    assert "\U0001F600" in masked
    assert "\U0001F389" in masked


def test_mixed_line_endings_crlf(tmp_path):
    patterns = build_patterns()
    log_file = tmp_path / "crlf.log"
    log_file.write_bytes(b"Contact john.doe@example.com\r\nNo secrets here\r\n")

    lines = list(read_log_file_lines(str(log_file)))
    masked_lines, total_counts, line_count = mask_log(lines, patterns)

    assert line_count == 2
    assert total_counts["email"] == 1


def test_very_long_line_does_not_crash():
    patterns = build_patterns()
    padding = "x" * 5000
    line = f"{padding} contact john.doe@example.com {padding}\n"
    masked, counts = mask_line(line, patterns)
    assert counts["email"] == 1
    assert "[EMAIL_REDACTED]" in masked


def test_main_check_mode_exits_1_when_secrets_found(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    log_file = tmp_path / "leaky.log"
    log_file.write_text("Contact john.doe@example.com\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(log_file), "--check"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 1
    assert "FAILED" in capsys.readouterr().out


def test_main_check_mode_exits_0_when_clean(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    log_file = tmp_path / "clean.log"
    log_file.write_text("Nothing sensitive in this line.\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(log_file), "--check"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 0
    assert "PASSED" in capsys.readouterr().out


def test_main_output_flag_rejected_with_multiple_files(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    file_a = tmp_path / "a.log"
    file_b = tmp_path / "b.log"
    file_a.write_text("Nothing sensitive.\n", encoding="utf-8")
    file_b.write_text("Nothing sensitive.\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(file_a), str(file_b), "-o", "out.log"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    assert "single input file" in capsys.readouterr().out


def test_main_output_flag_rejected_with_check_mode(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    log_file = tmp_path / "a.log"
    log_file.write_text("Nothing sensitive.\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(log_file), "--check", "-o", "out.log"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    assert "--check" in capsys.readouterr().out


def test_expand_filepaths_passes_through_plain_files(tmp_path):
    log_file = tmp_path / "a.log"
    log_file.write_text("hello\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(log_file)], ["log"], recursive=True)

    assert expanded == [str(log_file)]
    assert warnings == []


def test_expand_filepaths_expands_directory_recursive(tmp_path):
    (tmp_path / "a.log").write_text("top level\n", encoding="utf-8")
    nested_dir = tmp_path / "nested"
    nested_dir.mkdir()
    (nested_dir / "b.log").write_text("nested\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], ["log"], recursive=True)

    assert len(expanded) == 2
    assert any("a.log" in p for p in expanded)
    assert any("b.log" in p for p in expanded)
    assert warnings == []


def test_expand_filepaths_respects_non_recursive(tmp_path):
    (tmp_path / "a.log").write_text("top level\n", encoding="utf-8")
    nested_dir = tmp_path / "nested"
    nested_dir.mkdir()
    (nested_dir / "b.log").write_text("nested\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], ["log"], recursive=False)

    assert len(expanded) == 1
    assert "a.log" in expanded[0]


def test_expand_filepaths_filters_by_extension(tmp_path):
    (tmp_path / "a.log").write_text("log file\n", encoding="utf-8")
    (tmp_path / "b.txt").write_text("txt file\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], ["log"], recursive=True)

    assert len(expanded) == 1
    assert "a.log" in expanded[0]


def test_expand_filepaths_multiple_extensions(tmp_path):
    (tmp_path / "a.log").write_text("log file\n", encoding="utf-8")
    (tmp_path / "b.txt").write_text("txt file\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], ["log", "txt"], recursive=True)

    assert len(expanded) == 2


def test_expand_filepaths_warns_on_empty_directory(tmp_path):
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()

    expanded, warnings = expand_filepaths([str(empty_dir)], ["log"], recursive=True)

    assert expanded == []
    assert len(warnings) == 1
    assert "empty" in warnings[0]


def test_load_config_valid(tmp_path):
    config_file = tmp_path / ".log-guard.json"
    config_file.write_text('{"extensions": ["log", "txt"], "recursive": false, "quiet": true}', encoding="utf-8")

    config = load_config(str(config_file))
    assert config["extensions"] == ["log", "txt"]
    assert config["recursive"] is False
    assert config["quiet"] is True


def test_load_config_rejects_non_dict(tmp_path):
    config_file = tmp_path / ".log-guard.json"
    config_file.write_text('["log", "txt"]', encoding="utf-8")

    with pytest.raises(ValueError):
        load_config(str(config_file))


def test_load_config_rejects_unknown_key(tmp_path):
    config_file = tmp_path / ".log-guard.json"
    config_file.write_text('{"made_up_option": true}', encoding="utf-8")

    with pytest.raises(ValueError):
        load_config(str(config_file))


def test_load_config_rejects_bad_extensions_type(tmp_path):
    config_file = tmp_path / ".log-guard.json"
    config_file.write_text('{"extensions": "log"}', encoding="utf-8")

    with pytest.raises(ValueError):
        load_config(str(config_file))


def test_load_config_rejects_bad_recursive_type(tmp_path):
    config_file = tmp_path / ".log-guard.json"
    config_file.write_text('{"recursive": "yes"}', encoding="utf-8")

    with pytest.raises(ValueError):
        load_config(str(config_file))


def test_expand_filepaths_ext_all_matches_every_file(tmp_path):
    (tmp_path / "a.log").write_text("log\n", encoding="utf-8")
    (tmp_path / "b.dat").write_text("dat\n", encoding="utf-8")
    (tmp_path / "c.xyz").write_text("xyz\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], None, recursive=True)

    assert len(expanded) == 3
    assert warnings == []


def test_expand_filepaths_ext_all_empty_directory_warning(tmp_path):
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()

    expanded, warnings = expand_filepaths([str(empty_dir)], None, recursive=True)

    assert expanded == []
    assert len(warnings) == 1
    assert "extensions" not in warnings[0].lower()  # generic message, not extension-specific


def test_expand_filepaths_skips_ignored_directories(tmp_path):
    (tmp_path / "app.log").write_text("real log\n", encoding="utf-8")

    node_modules = tmp_path / "node_modules"
    node_modules.mkdir()
    (node_modules / "dependency.log").write_text("should be skipped\n", encoding="utf-8")

    git_dir = tmp_path / ".git"
    git_dir.mkdir()
    (git_dir / "config.log").write_text("should be skipped\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(tmp_path)], ["log"], recursive=True)

    assert len(expanded) == 1
    assert "app.log" in expanded[0]
    assert not any("node_modules" in p for p in expanded)
    assert not any(".git" in p for p in expanded)


def test_expand_filepaths_ignored_dirs_dont_apply_when_pointed_at_directly(tmp_path):
    """
    If the user explicitly passes an ignored-name directory as their
    input, it should still be scanned — the ignore-list only applies
    while auto-recursing into subfolders.
    """
    venv_dir = tmp_path / "venv"
    venv_dir.mkdir()
    (venv_dir / "info.log").write_text("explicitly targeted\n", encoding="utf-8")

    expanded, warnings = expand_filepaths([str(venv_dir)], ["log"], recursive=True)

    assert len(expanded) == 1
    assert "info.log" in expanded[0]


def test_main_ext_all_scans_nonstandard_extensions(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data.dat").write_text("Contact john.doe@example.com\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(tmp_path), "--check", "--ext", "all"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 1
    assert "FAILED" in capsys.readouterr().out


def test_main_directory_scanning_check_mode(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "a.log").write_text("Contact john.doe@example.com\n", encoding="utf-8")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "b.log").write_text("Call 555-123-4567\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(tmp_path), "--check"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 1
    output = capsys.readouterr().out
    assert "FAILED" in output
    assert "found 2 potential secret(s) across 2 file(s)" in output


def test_main_recursive_and_no_recursive_conflict_errors(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "a.log").write_text("hello\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(tmp_path), "--recursive", "--no-recursive"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    assert "cannot use both" in capsys.readouterr().out


def test_main_empty_directory_errors(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    empty_dir = tmp_path / "empty"
    empty_dir.mkdir()

    monkeypatch.setattr(sys, "argv", ["log-guard", str(empty_dir)])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    output = capsys.readouterr().out
    assert "no files matching" in output
    assert "no files to scan" in output


# ---------------------------------------------------------------------------
# .xlsx spreadsheet support
# ---------------------------------------------------------------------------


def _make_xlsx(path, sheets, hidden=()):
    """
    Builds a small .xlsx file for a test and returns its path as a string.

    `sheets` maps each sheet name to a dict of {cell coordinate: value},
    e.g. {"Users": {"A1": "alice@example.com", "B1": 5551234567}}.
    Sheet names listed in `hidden` are saved as hidden sheets.
    """
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, cells in sheets.items():
        worksheet = workbook.create_sheet(title)
        for coordinate, value in cells.items():
            worksheet[coordinate] = value
        if title in hidden:
            worksheet.sheet_state = "hidden"
    workbook.save(str(path))
    return str(path)


def _read_cell(path, sheet_title, coordinate):
    """
    Opens a saved .xlsx file and returns one cell's value.
    """
    workbook = openpyxl.load_workbook(str(path))
    return workbook[sheet_title][coordinate].value


def test_mask_workbook_masks_text_cells(tmp_path):
    path = _make_xlsx(
        tmp_path / "data.xlsx",
        {"Sheet1": {"A1": "Contact john.doe@example.com", "A2": "Nothing sensitive"}},
    )

    workbook, counts, cell_count, sheet_count = mask_workbook(path, build_patterns())

    assert counts["email"] == 1
    assert cell_count == 2
    assert sheet_count == 1
    assert workbook["Sheet1"]["A1"].value == "Contact [EMAIL_REDACTED]"
    assert workbook["Sheet1"]["A2"].value == "Nothing sensitive"


def test_mask_workbook_scans_whole_number_cells(tmp_path):
    path = _make_xlsx(tmp_path / "data.xlsx", {"Sheet1": {"A1": 5551234567}})

    workbook, counts, cell_count, sheet_count = mask_workbook(path, build_patterns())

    assert counts["phone"] == 1
    assert workbook["Sheet1"]["A1"].value == "[PHONE_REDACTED]"


def test_mask_workbook_skips_formulas(tmp_path):
    formula = '=HYPERLINK("mailto:support@example.com","Email support")'
    path = _make_xlsx(tmp_path / "data.xlsx", {"Sheet1": {"A1": formula}})

    workbook, counts, cell_count, sheet_count = mask_workbook(path, build_patterns())

    assert counts["email"] == 0
    assert cell_count == 1
    assert workbook["Sheet1"]["A1"].value == formula


def test_mask_workbook_skips_dates_decimals_and_booleans(tmp_path):
    path = _make_xlsx(
        tmp_path / "data.xlsx",
        {"Sheet1": {"A1": datetime(2024, 1, 15), "A2": 3.75, "A3": True}},
    )

    workbook, counts, cell_count, sheet_count = mask_workbook(path, build_patterns())

    assert cell_count == 3
    assert all(count == 0 for count in counts.values())
    assert workbook["Sheet1"]["A1"].value == datetime(2024, 1, 15)
    assert workbook["Sheet1"]["A2"].value == 3.75
    assert workbook["Sheet1"]["A3"].value is True


def test_mask_workbook_scans_hidden_sheets(tmp_path):
    path = _make_xlsx(
        tmp_path / "data.xlsx",
        {
            "Visible": {"A1": "Nothing here"},
            "Secret": {"A1": "key sk_live_51Hz8f92jak3ndlka9d"},
        },
        hidden=("Secret",),
    )

    workbook, counts, cell_count, sheet_count = mask_workbook(path, build_patterns())

    assert sheet_count == 2
    assert counts["api_key"] == 1
    assert workbook["Secret"]["A1"].value == "key [API_KEY_REDACTED]"


def test_mask_workbook_applies_custom_patterns(tmp_path):
    patterns_file = tmp_path / "custom_patterns.json"
    patterns_file.write_text(r'{"employee_id": "EMP-\\d{6}"}', encoding="utf-8")
    patterns = build_patterns(str(patterns_file))
    path = _make_xlsx(tmp_path / "data.xlsx", {"Staff": {"A1": "Badge EMP-482913"}})

    workbook, counts, cell_count, sheet_count = mask_workbook(path, patterns)

    assert counts["employee_id"] == 1
    assert workbook["Staff"]["A1"].value == "Badge [EMPLOYEE_ID_REDACTED]"


def test_process_file_xlsx_writes_masked_copy_and_keeps_original(tmp_path):
    path = _make_xlsx(tmp_path / "data.xlsx", {"Users": {"A1": "alice@example.com"}})

    result = process_file(path, build_patterns(), check_mode=False)

    assert result["error"] is None
    assert result["kind"] == "spreadsheet"
    assert result["cell_count"] == 1
    assert result["sheet_count"] == 1
    output_path = tmp_path / "data.masked.xlsx"
    assert result["output_filepath"] == str(output_path)
    assert _read_cell(output_path, "Users", "A1") == "[EMAIL_REDACTED]"
    assert _read_cell(path, "Users", "A1") == "alice@example.com"


def test_process_file_xlsx_check_mode_writes_nothing(tmp_path):
    path = _make_xlsx(tmp_path / "data.xlsx", {"Users": {"A1": "alice@example.com"}})

    result = process_file(path, build_patterns(), check_mode=True)

    assert result["counts"]["email"] == 1
    assert result["output_filepath"] is None
    assert not (tmp_path / "data.masked.xlsx").exists()


def test_process_file_xlsx_clean_file_skips_write(tmp_path):
    path = _make_xlsx(tmp_path / "clean.xlsx", {"Sheet1": {"A1": "Nothing sensitive"}})

    result = process_file(path, build_patterns(), check_mode=False)

    assert result["skipped_write"] is True
    assert result["output_filepath"] is None
    assert not (tmp_path / "clean.masked.xlsx").exists()


def test_process_file_xlsx_empty_workbook(tmp_path):
    path = str(tmp_path / "empty.xlsx")
    openpyxl.Workbook().save(path)

    result = process_file(path, build_patterns(), check_mode=True)

    assert result["error"] is None
    assert result["empty"] is True
    assert result["cell_count"] == 0


def test_process_file_corrupt_xlsx_reports_clear_error(tmp_path):
    bad_file = tmp_path / "broken.xlsx"
    bad_file.write_text("This is not really a spreadsheet.\n", encoding="utf-8")

    result = process_file(str(bad_file), build_patterns(), check_mode=True)

    assert result["error"] is not None
    assert "could not read spreadsheet" in result["error"]


def test_process_file_xls_reports_not_supported(tmp_path):
    old_file = tmp_path / "old.xls"
    old_file.write_text("Pretend old Excel file\n", encoding="utf-8")

    result = process_file(str(old_file), build_patterns(), check_mode=True)

    assert result["error"] is not None
    assert "not supported" in result["error"]
    assert ".xlsx" in result["error"]


def test_process_file_xlsx_extension_is_case_insensitive(tmp_path):
    path = _make_xlsx(tmp_path / "REPORT.XLSX", {"Sheet1": {"A1": "alice@example.com"}})

    result = process_file(path, build_patterns(), check_mode=True)

    assert result["error"] is None
    assert result["kind"] == "spreadsheet"
    assert result["counts"]["email"] == 1


def test_default_extensions_include_xlsx_but_not_xls():
    assert "xlsx" in DEFAULT_EXTENSIONS
    assert "xls" not in DEFAULT_EXTENSIONS


def test_main_xlsx_normal_mode_reports_cells_and_writes_copy(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    path = _make_xlsx(
        tmp_path / "data.xlsx",
        {"Sheet1": {"A1": "john.doe@example.com", "B1": "Call 555-123-4567"}},
    )

    monkeypatch.setattr(sys, "argv", ["log-guard", path])
    main()

    output = capsys.readouterr().out
    assert "Scanned 2 cells across 1 sheet(s)" in output
    assert "Masked spreadsheet written to" in output
    assert (tmp_path / "data.masked.xlsx").exists()


def test_main_directory_scan_includes_xlsx_by_default(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "a.log").write_text("Contact john.doe@example.com\n", encoding="utf-8")
    nested = tmp_path / "nested"
    nested.mkdir()
    _make_xlsx(nested / "report.xlsx", {"Sheet1": {"A1": "Call 555-123-4567"}})

    monkeypatch.setattr(sys, "argv", ["log-guard", str(tmp_path), "--check"])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 1
    assert "found 2 potential secret(s) across 2 file(s)" in capsys.readouterr().out


def test_main_xls_file_exits_with_error(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    old_file = tmp_path / "old.xls"
    old_file.write_text("Pretend old Excel file\n", encoding="utf-8")

    monkeypatch.setattr(sys, "argv", ["log-guard", str(old_file)])

    with pytest.raises(SystemExit) as exc_info:
        main()

    assert exc_info.value.code == 2
    assert "not supported" in capsys.readouterr().out